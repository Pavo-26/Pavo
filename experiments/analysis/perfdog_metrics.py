#!/usr/bin/env python3
"""Calculate weighted jank and average power from PerfDog XLSX exports."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
import warnings
from collections import defaultdict
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

FIELDS = {
    "SmallJank(/10min)": "small_jank_per_10min",
    "Jank(/10min)": "jank_per_10min",
    "BigJank(/10min)": "big_jank_per_10min",
    "Avg(Power)[mW]": "avg_power_mw",
}


def read_metrics(path: Path) -> dict[str, float]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl: python -m pip install openpyxl==3.1.5") from exc

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore", message="Workbook contains no default style", category=UserWarning,
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "all" not in workbook.sheetnames:
            raise ValueError("Missing PerfDog worksheet 'all'")
        sheet = workbook["all"]
        # Some PerfDog exports incorrectly declare their used range as A1:A1.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=40, values_only=True)
        for row in rows:
            if row and str(row[0]).strip() == "Stat":
                headers = [str(value).strip() for value in next(rows, ())]
                values = next(rows, ())
                break
        else:
            raise ValueError("Missing PerfDog Stat table")

        metrics = {}
        for header, name in FIELDS.items():
            if headers.count(header) != 1:
                raise ValueError(f"Expected one Stat column {header!r}")
            column = headers.index(header)
            value = values[column] if column < len(values) else None
            if (isinstance(value, bool) or not isinstance(value, (int, float))
                    or not math.isfinite(value) or value < 0):
                raise ValueError(f"Invalid Stat value for {header!r}: {value!r}")
            metrics[name] = float(value)
        if metrics["avg_power_mw"] <= 0:
            raise ValueError("Avg(Power)[mW] must be positive")
        metrics["wj_per_10min"] = (
            metrics["small_jank_per_10min"]
            + 2 * metrics["jank_per_10min"]
            + 3 * metrics["big_jank_per_10min"]
        )
        return metrics
    finally:
        workbook.close()


def resolve_export(value: str, inventory: Path | None) -> Path:
    path = Path(value).expanduser()
    candidates = {path.resolve()}
    if inventory is not None and not path.is_absolute():
        candidates.add((inventory.parent / path).resolve())
    existing = [candidate for candidate in candidates if candidate.is_file()]
    if not existing:
        raise ValueError(f"Export not found: {value}")
    if len(existing) != 1:
        raise ValueError(f"Ambiguous export path: {value}; use an absolute path")
    path = existing[0]
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Expected an XLSX export: {path}")
    return path


def read_runs(inputs: list[Path], app: str | None, policy: str | None) -> list[dict]:
    inventory = None
    if len(inputs) == 1 and inputs[0].suffix.lower() == ".json":
        if app is not None or policy is not None:
            raise ValueError("App and policy are read from run_inventory.json")
        inventory = inputs[0].resolve()
        records = json.loads(inventory.read_text(encoding="utf-8-sig"))
        if not isinstance(records, list) or not records:
            raise ValueError("Expected a nonempty run_inventory.json list")
        for row in records:
            if not isinstance(row, dict) or row.get("status") != "valid":
                raise ValueError("Inventory contains an unsuccessful capture")
    else:
        if not app or not policy:
            raise ValueError("XLSX inputs require --app and --policy")
        records = [
            {"app": app, "policy": policy, "repetition": number, "copied_xlsx": str(path)}
            for number, path in enumerate(inputs, 1)
        ]

    runs, seen_runs, seen_files = [], set(), set()
    for row in records:
        app, policy = row.get("app"), row.get("policy")
        if any(not isinstance(value, str) or not value.strip() for value in (app, policy)):
            raise ValueError("Each run needs nonempty app and policy labels")
        repetition = row.get("repetition")
        if type(repetition) is not int or repetition < 1:
            raise ValueError("Each repetition must be a positive integer")
        value = row.get("copied_xlsx")
        if not isinstance(value, str) or not value.strip():
            raise ValueError("Each run needs copied_xlsx")
        path = resolve_export(value, inventory)
        key = (app.strip(), policy.strip(), repetition)
        if key in seen_runs:
            raise ValueError(f"Duplicate run: {key}")
        if path in seen_files:
            raise ValueError(f"Export listed more than once: {path}")
        seen_runs.add(key)
        seen_files.add(path)
        runs.append({
            "app": key[0], "policy": key[1], "repetition": repetition, "file": str(path),
        })
    return sorted(runs, key=lambda row: (row["app"], row["policy"], row["repetition"]))


def summarize(runs: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in runs:
        groups[(row["app"], row["policy"])].append(row)
    summaries = []
    for (app, policy), group in sorted(groups.items()):
        row = {"app": app, "policy": policy, "n": len(group)}
        for source, prefix, unit in (
            ("wj_per_10min", "wj", "per_10min"),
            ("avg_power_mw", "power", "mw"),
        ):
            values = [run[source] for run in group]
            row[f"{prefix}_mean_{unit}"] = statistics.mean(values)
            row[f"{prefix}_sd_{unit}"] = statistics.stdev(values) if len(values) > 1 else ""
        summaries.append(row)
    return summaries


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", type=Path,
                        help="One run_inventory.json, or XLSX files from one App/policy")
    parser.add_argument("--app", help="App label for direct XLSX inputs")
    parser.add_argument("--policy", help="Policy label for direct XLSX inputs")
    parser.add_argument("--output-dir", type=Path, required=True,
                        help="New directory for runs.csv and summary.csv")
    args = parser.parse_args(argv)
    try:
        runs = read_runs(args.inputs, args.app, args.policy)
        for run in runs:
            try:
                run.update(read_metrics(Path(run["file"])))
            except (OSError, ValueError, RuntimeError, BadZipFile, ParseError) as exc:
                raise ValueError(f"{run['file']}: {exc}") from exc
        summaries = summarize(runs)
        args.output_dir.mkdir(parents=True, exist_ok=False)
        write_csv(args.output_dir / "runs.csv", runs)
        write_csv(args.output_dir / "summary.csv", summaries)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    print(args.output_dir / "runs.csv")
    print(args.output_dir / "summary.csv")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
